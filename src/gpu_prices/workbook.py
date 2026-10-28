from __future__ import annotations

import csv
import posixpath
import re
from decimal import Decimal
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from gpu_prices.models import Observation, SourceRow

SOURCE_SHEET = "算力价格来源"
EXPECTED_HEADERS = [
    "供应商名称",
    "供应商类型",
    "卡型",
    "价格类型",
    "地区",
    "云类型",
    "租期（月）",
    "价格链接",
    "截图",
    "备注",
]
OUTPUT_HEADERS = [
    "dtstatdate",
    "supplier_name",
    "supplier_type",
    "gpu_model_code",
    "price_type",
    "region",
    "cloud_type",
    "term",
    "url",
    "price",
    "data_source",
]
DATA_SOURCE = "price_mining"
AUDIT_HEADERS = [
    "workbook_row",
    *OUTPUT_HEADERS,
    "status",
    "effective_at",
    "fetched_at_utc",
    "source_url",
    "raw_value",
    "note",
]


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CELL_COLUMN_RE = re.compile(r"([A-Z]+)")


def _xml_sheet_rows(path: Path, sheet_name: str) -> list[tuple[int, dict[str, str]]]:
    """Read cell values without parsing styles.

    The source workbook contains empty Excel fill records that openpyxl rejects,
    although its worksheet data is valid. Reading the OOXML values directly
    keeps ingestion independent of workbook styling.
    """
    with ZipFile(path) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in root.findall(f"{{{_MAIN_NS}}}si"):
                shared_strings.append(
                    "".join(node.text or "" for node in item.iter(f"{{{_MAIN_NS}}}t"))
                )

        workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relationships = ElementTree.fromstring(
            archive.read("xl/_rels/workbook.xml.rels")
        )
        relationship_targets = {
            node.attrib["Id"]: node.attrib["Target"]
            for node in relationships.findall(f"{{{_PACKAGE_REL_NS}}}Relationship")
        }
        sheet_node = next(
            (
                node
                for node in workbook.findall(f".//{{{_MAIN_NS}}}sheet")
                if node.attrib["name"] == sheet_name
            ),
            None,
        )
        if sheet_node is None:
            raise ValueError(f"Workbook is missing sheet {sheet_name!r}")
        relationship_id = sheet_node.attrib[f"{{{_OFFICE_REL_NS}}}id"]
        sheet_path = posixpath.normpath(
            posixpath.join("xl", relationship_targets[relationship_id])
        )
        sheet = ElementTree.fromstring(archive.read(sheet_path))

        rows: list[tuple[int, dict[str, str]]] = []
        for row_node in sheet.findall(f".//{{{_MAIN_NS}}}sheetData/{{{_MAIN_NS}}}row"):
            row_number = int(row_node.attrib["r"])
            values: dict[str, str] = {}
            for cell in row_node.findall(f"{{{_MAIN_NS}}}c"):
                reference = cell.attrib["r"]
                match = _CELL_COLUMN_RE.match(reference)
                if match is None:
                    continue
                column = match.group(1)
                cell_type = cell.attrib.get("t")
                value_node = cell.find(f"{{{_MAIN_NS}}}v")
                if cell_type == "inlineStr":
                    value = "".join(
                        node.text or "" for node in cell.iter(f"{{{_MAIN_NS}}}t")
                    )
                elif value_node is None or value_node.text is None:
                    value = ""
                elif cell_type == "s":
                    value = shared_strings[int(value_node.text)]
                elif cell_type == "b":
                    value = "TRUE" if value_node.text == "1" else "FALSE"
                else:
                    value = value_node.text
                values[column] = value
            rows.append((row_number, values))
        return rows


def _parse_term(value: str, row_number: int) -> Decimal | None:
    if not value or value.upper() == "N/A":
        return None
    try:
        return Decimal(value)
    except Exception as exc:
        raise ValueError(
            f"Invalid term {value!r} at workbook row {row_number}"
        ) from exc


def load_source_rows(path: Path) -> list[SourceRow]:
    xml_rows = _xml_sheet_rows(path, SOURCE_SHEET)
    first_row = next(
        (values for row_number, values in xml_rows if row_number == 1),
        {},
    )
    headers = [
        _clean(first_row.get(get_column_letter(index), ""))
        for index in range(1, len(EXPECTED_HEADERS) + 1)
    ]
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"Unexpected {SOURCE_SHEET!r} headers: {headers}")

    rows: list[SourceRow] = []
    for row_number, cells in xml_rows:
        if row_number < 2:
            continue
        values = [
            _clean(cells.get(get_column_letter(index), "")) for index in range(1, 9)
        ]
        if not any(values):
            continue
        required = values[:6] + [values[7]]
        if not all(required):
            raise ValueError(
                f"Incomplete source definition at workbook row {row_number}"
            )
        gpu_model = values[2]
        if values[0] == "TheAI Cloud" and gpu_model == "H200":
            # The workbook screenshots and the supplier API both identify B300;
            # TheAI discontinued H200 before this data range.
            gpu_model = "B300"
        rows.append(
            SourceRow(
                workbook_row=row_number,
                source_name=values[0],
                supplier_type=values[1],
                gpu_model_code=gpu_model,
                price_type=values[3],
                region=values[4],
                cloud_type=values[5],
                term=_parse_term(values[6], row_number),
                url=values[7],
            )
        )
    return rows


def _price_text(observation: Observation) -> str:
    return "" if observation.price is None else str(observation.price)


def _data_row(run_date: str, observation: Observation) -> list[object]:
    row = observation.row
    return [
        run_date,
        row.source_name,
        row.supplier_type,
        row.gpu_model_code,
        row.price_type,
        row.region,
        row.cloud_type,
        "" if row.term is None else str(row.term),
        row.url,
        _price_text(observation),
        DATA_SOURCE,
    ]


def write_outputs(
    observations: Iterable[Observation],
    *,
    run_date: str,
    output_dir: Path,
) -> tuple[Path, Path, Path]:
    ordered = sorted(observations, key=lambda item: item.row.workbook_row)
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / f"gpu_prices_{run_date}.csv"
    xlsx_path = output_dir / f"gpu_prices_{run_date}.xlsx"
    audit_path = output_dir / f"gpu_prices_{run_date}_audit.csv"

    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(OUTPUT_HEADERS)
        for observation in ordered:
            writer.writerow(_data_row(run_date, observation))

    with audit_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(AUDIT_HEADERS)
        for observation in ordered:
            writer.writerow(
                [
                    observation.row.workbook_row,
                    *_data_row(run_date, observation),
                    observation.status,
                    observation.effective_at,
                    observation.fetched_at.isoformat(),
                    observation.source_url,
                    observation.raw_value,
                    observation.note,
                ]
            )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "每日价格"
    sheet.append(OUTPUT_HEADERS)
    for observation in ordered:
        row = _data_row(run_date, observation)
        row[7] = (
            float(observation.row.term) if observation.row.term is not None else None
        )
        row[9] = float(observation.price) if observation.price is not None else None
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    widths = [14, 22, 16, 18, 16, 20, 16, 12, 48, 14, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    table = Table(displayName="DailyGpuPrices", ref=f"A1:K{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    workbook.save(xlsx_path)
    return csv_path, xlsx_path, audit_path
